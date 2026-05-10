#!/usr/bin/env python3
"""
build_sale_v2.py — Rosey's Resource Sale
Reads the CSV and builds the initial Excel file with formatting.
Run this first, then run fetch_images.py to populate cover images.

    pip install openpyxl
    python build_sale_v2.py
"""

import csv
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
CSV_FILE   = os.path.join(SCRIPT_DIR, "Rosey's Resource Sale - Sheet1.csv")
XLSX_FILE  = os.path.join(SCRIPT_DIR, "Rosey's Resource Sale.xlsx")

# Output columns (CSV col index -> output header)
# CSV: 0=Title(Mar), 1=Author, 2=empty(skip), 3=Date, 4=Description,
#      5=CD/DVD, 6=Cost, 7=Book Scan, 8=Audio or DVD, 9=Link 3
HEADERS = [
    "#",
    "Title",
    "Author",
    "Date of Publication",
    "Description",
    "CD / DVD / Online Access Code",
    "Cost I Paid",
    "Book Scan",
    "Audio or DVD",
    "Link 3",
    "Image Filename",
    "Image Link",
]

# Approximate column widths
COL_WIDTHS = {
    "#": 5,
    "Title": 45,
    "Author": 30,
    "Date of Publication": 12,
    "Description": 22,
    "CD / DVD / Online Access Code": 18,
    "Cost I Paid": 10,
    "Book Scan": 14,
    "Audio or DVD": 14,
    "Link 3": 14,
    "Image Filename": 30,
    "Image Link": 12,
}

URL_COLS = {"Book Scan", "Audio or DVD", "Link 3"}

HEADER_FILL = PatternFill("solid", fgColor="BDD7EE")
HEADER_FONT = Font(name="Arial", size=10, bold=True)
CELL_FONT   = Font(name="Arial", size=10)
LINK_FONT   = Font(name="Arial", size=10, color="0563C1", underline="single")

with open(CSV_FILE, "r", encoding="utf-8") as f:
    rows = list(csv.reader(f))

data_rows = rows[1:]  # skip header
print(f"Read {len(data_rows)} rows from CSV")

wb = Workbook()
ws = wb.active
ws.title = "Sale Items"

# Write header
for col_idx, header in enumerate(HEADERS, 1):
    cell = ws.cell(row=1, column=col_idx, value=header)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

ws.row_dimensions[1].height = 30
ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}1"

# Write data
for row_num, row in enumerate(data_rows, 1):
    while len(row) < 10:
        row.append("")

    # Map CSV columns (skip col index 2 which is empty)
    title   = row[0].strip()
    author  = row[1].strip()
    date    = row[3].strip()
    desc    = row[4].strip()
    media   = row[5].strip()
    cost    = row[6].strip()
    scan    = row[7].strip()
    audio   = row[8].strip()
    link3   = row[9].strip()

    excel_row = row_num + 1
    values = [row_num, title, author, date, desc, media, cost, scan, audio, link3, "", ""]

    for col_idx, (header, value) in enumerate(zip(HEADERS, values), 1):
        cell = ws.cell(row=excel_row, column=col_idx)

        if header == "#":
            cell.value = value  # integer
        elif header in URL_COLS and value:
            cell.value = "Link"
            cell.hyperlink = value
            cell.font = LINK_FONT
            continue
        else:
            cell.value = value

        cell.font = CELL_FONT
        cell.alignment = Alignment(vertical="top", wrap_text=(header == "Title"))

# Set column widths
for col_idx, header in enumerate(HEADERS, 1):
    ws.column_dimensions[get_column_letter(col_idx)].width = COL_WIDTHS.get(header, 15)

# Freeze header row
ws.freeze_panes = "A2"

wb.save(XLSX_FILE)
print(f"Saved: {XLSX_FILE}")
print(f"Rows written: {len(data_rows)}")
print("Done! Now run fetch_images.py to download cover images.")
